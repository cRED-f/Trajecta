"""Local document, archive, image and OCR tools."""

from __future__ import annotations

import csv
import hashlib
import json
import mimetypes
import os
import shutil
import tarfile
import zipfile
from pathlib import Path
from typing import Any

import yaml

from server.src.config import Settings


class VirtualPathResolver:
    """Map agent virtual paths to explicitly exposed host roots."""

    def __init__(
        self,
        settings_or_workspace: Settings | str,
        uploads_root: str | None = None,
    ) -> None:
        if isinstance(settings_or_workspace, Settings):
            workspace = settings_or_workspace.tools.workspace_root
            uploads = settings_or_workspace.chat.uploads_path
        else:
            if uploads_root is None:
                raise ValueError("uploads_root is required when passing a workspace path")
            workspace = settings_or_workspace
            uploads = uploads_root
        self._workspace = Path(workspace).resolve()
        self._uploads = Path(uploads).resolve()

    @property
    def workspace_root(self) -> Path:
        return self._workspace

    @property
    def uploads_root(self) -> Path:
        return self._uploads

    @staticmethod
    def _inside(root: Path, candidate: Path) -> Path:
        resolved = candidate.resolve()
        try:
            resolved.relative_to(root)
        except ValueError as exc:
            raise ValueError("Path escapes the allowed root") from exc
        return resolved

    def resolve(self, virtual_path: str, *, writable: bool = False) -> Path:
        if virtual_path == "/workspace" or virtual_path.startswith("/workspace/"):
            rel = virtual_path[len("/workspace/") :] if virtual_path != "/workspace" else ""
            return self._inside(self._workspace, self._workspace / rel)
        if virtual_path == "/uploads" or virtual_path.startswith("/uploads/"):
            if writable:
                raise PermissionError("/uploads is read-only")
            rel = virtual_path[len("/uploads/") :] if virtual_path != "/uploads" else ""
            return self._inside(self._uploads, self._uploads / rel)
        raise ValueError("Use /workspace/... or /uploads/... virtual paths")

    def virtualize(self, path: Path) -> str:
        resolved = path.resolve()
        try:
            return "/workspace/" + resolved.relative_to(self._workspace).as_posix()
        except ValueError:
            pass
        try:
            return "/uploads/" + resolved.relative_to(self._uploads).as_posix()
        except ValueError as exc:
            raise ValueError("Path is outside exposed roots") from exc


class DocumentTools:
    def __init__(self, settings: Settings, workspace_root: str | None = None) -> None:
        # `workspace_root` rebinds /workspace/ to one conversation's folder
        # without touching the shared settings object.
        if workspace_root is None:
            self.paths = VirtualPathResolver(settings)
        else:
            self.paths = VirtualPathResolver(workspace_root, settings.chat.uploads_path)

    @staticmethod
    def _bounded(text: str, max_chars: int) -> str:
        return text[: max(1, min(int(max_chars), 200_000))]

    def metadata(self, virtual_path: str) -> dict[str, Any]:
        path = self.paths.resolve(virtual_path)
        if not path.exists():
            raise FileNotFoundError(virtual_path)
        stat = path.stat()
        result: dict[str, Any] = {
            "path": virtual_path,
            "name": path.name,
            "is_dir": path.is_dir(),
            "size_bytes": stat.st_size,
            "modified_at": stat.st_mtime,
            "mime_type": mimetypes.guess_type(path.name)[0],
        }
        if path.is_file():
            digest = hashlib.sha256()
            with path.open("rb") as handle:
                for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                    digest.update(chunk)
            result["sha256"] = digest.hexdigest()
        return result

    def read(self, virtual_path: str, *, max_chars: int = 60_000) -> dict[str, Any]:
        path = self.paths.resolve(virtual_path)
        if not path.is_file():
            raise FileNotFoundError(virtual_path)
        suffix = path.suffix.lower()

        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            pages = []
            for index, page in enumerate(reader.pages):
                pages.append({"page": index + 1, "text": page.extract_text() or ""})
            text = "\n\n".join(f"[Page {p['page']}]\n{p['text']}" for p in pages)
            return {"path": virtual_path, "kind": "pdf", "pages": len(pages), "text": self._bounded(text, max_chars)}

        if suffix == ".docx":
            from docx import Document

            doc = Document(str(path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            tables: list[list[list[str]]] = []
            for table in doc.tables:
                tables.append([[cell.text for cell in row.cells] for row in table.rows])
            return {
                "path": virtual_path,
                "kind": "docx",
                "text": self._bounded("\n".join(paragraphs), max_chars),
                "tables": tables[:50],
            }

        if suffix in {".xlsx", ".xlsm"}:
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise RuntimeError("openpyxl is required to read XLSX files") from exc
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheets: dict[str, list[list[Any]]] = {}
            total_chars = 0
            for sheet in workbook.worksheets:
                rows: list[list[Any]] = []
                for row in sheet.iter_rows(values_only=True):
                    values = list(row)
                    rows.append(values)
                    total_chars += sum(len(str(v)) for v in values if v is not None)
                    if total_chars >= max_chars:
                        break
                sheets[sheet.title] = rows
                if total_chars >= max_chars:
                    break
            return {"path": virtual_path, "kind": "xlsx", "sheets": sheets, "truncated": total_chars >= max_chars}

        if suffix == ".pptx":
            try:
                from pptx import Presentation
            except ImportError as exc:
                raise RuntimeError("python-pptx is required to read PPTX files") from exc
            presentation = Presentation(str(path))
            slides: list[dict[str, Any]] = []
            for index, slide in enumerate(presentation.slides):
                texts: list[str] = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        texts.append(str(shape.text))
                slides.append({"slide": index + 1, "text": "\n".join(texts)})
            text = "\n\n".join(f"[Slide {s['slide']}]\n{s['text']}" for s in slides)
            return {"path": virtual_path, "kind": "pptx", "slides": len(slides), "text": self._bounded(text, max_chars)}

        if suffix == ".csv":
            rows: list[list[str]] = []
            with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
                reader = csv.reader(handle)
                char_count = 0
                for row in reader:
                    rows.append(row)
                    char_count += sum(map(len, row))
                    if char_count >= max_chars:
                        break
            return {"path": virtual_path, "kind": "csv", "rows": rows, "truncated": char_count >= max_chars}

        if suffix == ".json":
            data = json.loads(path.read_text(encoding="utf-8"))
            rendered = json.dumps(data, ensure_ascii=False, indent=2)
            return {"path": virtual_path, "kind": "json", "text": self._bounded(rendered, max_chars)}

        if suffix in {".yaml", ".yml"}:
            data = yaml.safe_load(path.read_text(encoding="utf-8"))
            rendered = yaml.safe_dump(data, allow_unicode=True, sort_keys=False)
            return {"path": virtual_path, "kind": "yaml", "text": self._bounded(rendered, max_chars)}

        if suffix == ".xml" or suffix == ".html" or suffix == ".htm":
            text = path.read_text(encoding="utf-8", errors="replace")
            if suffix in {".html", ".htm"}:
                try:
                    from bs4 import BeautifulSoup
                    text = BeautifulSoup(text, "html.parser").get_text("\n")
                except ImportError:
                    pass
            return {"path": virtual_path, "kind": suffix.lstrip("."), "text": self._bounded(text, max_chars)}

        text = path.read_text(encoding="utf-8", errors="replace")
        return {"path": virtual_path, "kind": "text", "text": self._bounded(text, max_chars)}

    def search(self, virtual_path: str, query: str, *, max_matches: int = 50) -> list[dict[str, Any]]:
        extracted = self.read(virtual_path, max_chars=200_000)
        if "text" in extracted:
            lines = str(extracted["text"]).splitlines()
        else:
            lines = json.dumps(extracted, ensure_ascii=False).splitlines()
        needle = query.casefold()
        matches: list[dict[str, Any]] = []
        for line_no, line in enumerate(lines, start=1):
            if needle in line.casefold():
                matches.append({"line": line_no, "text": line[:2_000]})
                if len(matches) >= max(1, min(max_matches, 200)):
                    break
        return matches

    def ocr_image(self, virtual_path: str) -> dict[str, Any]:
        path = self.paths.resolve(virtual_path)
        try:
            import pytesseract
            from PIL import Image
        except ImportError as exc:
            raise RuntimeError("pytesseract and Pillow are required for OCR") from exc
        text = pytesseract.image_to_string(Image.open(path))
        return {"path": virtual_path, "text": text}

    def image_metadata(self, virtual_path: str) -> dict[str, Any]:
        path = self.paths.resolve(virtual_path)
        try:
            from PIL import Image
        except ImportError as exc:
            raise RuntimeError("Pillow is required for image tools") from exc
        with Image.open(path) as image:
            return {
                "path": virtual_path,
                "format": image.format,
                "mode": image.mode,
                "width": image.width,
                "height": image.height,
                "frames": getattr(image, "n_frames", 1),
                "exif": {str(k): str(v) for k, v in image.getexif().items()},
            }

    def image_transform(
        self,
        virtual_path: str,
        output_path: str,
        *,
        width: int | None = None,
        height: int | None = None,
        crop: list[int] | None = None,
        format: str | None = None,
    ) -> dict[str, Any]:
        source = self.paths.resolve(virtual_path)
        destination = self.paths.resolve(output_path, writable=True)
        destination.parent.mkdir(parents=True, exist_ok=True)
        try:
            from PIL import Image
        except ImportError as exc:
            raise RuntimeError("Pillow is required for image tools") from exc
        with Image.open(source) as image:
            if crop is not None:
                if len(crop) != 4:
                    raise ValueError("crop must be [left, top, right, bottom]")
                image = image.crop(tuple(int(v) for v in crop))
            if width is not None or height is not None:
                target_w = int(width or image.width)
                target_h = int(height or image.height)
                image = image.resize((target_w, target_h))
            image.save(destination, format=format)
        return self.image_metadata(output_path)

    def extract_archive(self, virtual_path: str, destination: str) -> dict[str, Any]:
        source = self.paths.resolve(virtual_path)
        dest = self.paths.resolve(destination, writable=True)
        dest.mkdir(parents=True, exist_ok=True)
        extracted: list[str] = []

        def safe_target(name: str) -> Path:
            target = (dest / name).resolve()
            try:
                target.relative_to(dest)
            except ValueError as exc:
                raise ValueError(f"Unsafe archive path: {name}") from exc
            return target

        if zipfile.is_zipfile(source):
            with zipfile.ZipFile(source) as archive:
                for info in archive.infolist():
                    target = safe_target(info.filename)
                    if info.is_dir():
                        target.mkdir(parents=True, exist_ok=True)
                        continue
                    target.parent.mkdir(parents=True, exist_ok=True)
                    with archive.open(info) as src, target.open("wb") as out:
                        shutil.copyfileobj(src, out)
                    extracted.append(self.paths.virtualize(target))
        elif tarfile.is_tarfile(source):
            with tarfile.open(source, "r:*") as archive:
                for member in archive.getmembers():
                    target = safe_target(member.name)
                    if member.isdir():
                        target.mkdir(parents=True, exist_ok=True)
                        continue
                    if not member.isfile():
                        continue
                    target.parent.mkdir(parents=True, exist_ok=True)
                    src = archive.extractfile(member)
                    if src is not None:
                        with target.open("wb") as out:
                            shutil.copyfileobj(src, out)
                        extracted.append(self.paths.virtualize(target))
        else:
            raise ValueError("Unsupported archive format")
        return {"destination": destination, "files": extracted[:1_000], "count": len(extracted)}

    def create_archive(self, paths: list[str], output_path: str, *, format: str = "zip") -> dict[str, Any]:
        sources = [(virtual, self.paths.resolve(virtual)) for virtual in paths]
        destination = self.paths.resolve(output_path, writable=True)
        destination.parent.mkdir(parents=True, exist_ok=True)
        if format == "zip":
            with zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                for virtual, source in sources:
                    if source.is_dir():
                        for child in source.rglob("*"):
                            if child.is_file():
                                archive.write(child, arcname=child.relative_to(source.parent))
                    else:
                        archive.write(source, arcname=source.name)
        elif format in {"tar", "tar.gz", "tgz"}:
            mode = "w:gz" if format in {"tar.gz", "tgz"} else "w"
            with tarfile.open(destination, mode) as archive:
                for _virtual, source in sources:
                    archive.add(source, arcname=source.name)
        else:
            raise ValueError("format must be zip, tar, tar.gz, or tgz")
        return self.metadata(output_path)

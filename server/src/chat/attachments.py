from __future__ import annotations

import asyncio
import hashlib
import mimetypes
import re

from pathlib import Path

from fastapi import UploadFile
from docx import Document
from pypdf import PdfReader

from server.src.chat.models import (
    Attachment,
    AttachmentKind,
    AttachmentStatus,
)

from server.src.chat.repository import (
    ChatRepository,
)


_IMAGE_EXTENSIONS = {
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".webp",
    ".heic",
    ".heif",
}

_TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".rst",
    ".py",
    ".js",
    ".jsx",
    ".ts",
    ".tsx",
    ".json",
    ".yaml",
    ".yml",
    ".toml",
    ".xml",
    ".html",
    ".css",
    ".scss",
    ".sql",
    ".java",
    ".go",
    ".rs",
    ".cpp",
    ".c",
    ".h",
    ".hpp",
    ".sh",
    ".ps1",
    ".bat",
    ".log",
    ".csv",
}

_DOCUMENT_EXTENSIONS = {
    ".ppt",
    ".pptx",
}


class AttachmentError(RuntimeError):
    pass


def _safe_filename(value: str) -> str:
    value = Path(value).name.strip()

    value = re.sub(
        r"[\x00-\x1f<>:\"/\\|?*]",
        "_",
        value,
    )

    if not value:
        return "attachment"

    return value[:255]


def _kind_for(
    filename: str,
) -> AttachmentKind:
    suffix = Path(filename).suffix.lower()

    if suffix in _IMAGE_EXTENSIONS:
        return AttachmentKind.IMAGE

    if suffix == ".pdf":
        return AttachmentKind.PDF

    if suffix == ".docx":
        return AttachmentKind.DOCX

    if suffix in _TEXT_EXTENSIONS:
        return AttachmentKind.TEXT

    if suffix in _DOCUMENT_EXTENSIONS:
        return AttachmentKind.DOCUMENT

    return AttachmentKind.OTHER


class AttachmentService:
    def __init__(
        self,
        repository: ChatRepository,
        *,
        uploads_root: str,
        max_upload_mb: int,
        max_extracted_chars: int,
    ) -> None:
        self._repository = repository

        self._root = Path(
            uploads_root
        ).resolve()

        self._root.mkdir(
            parents=True,
            exist_ok=True,
        )

        self._max_bytes = (
            max_upload_mb * 1024 * 1024
        )

        self._max_extracted_chars = (
            max_extracted_chars
        )

    async def save(
        self,
        *,
        conversation_id: str,
        upload: UploadFile,
    ) -> Attachment:
        import uuid

        attachment_id = uuid.uuid4().hex

        filename = _safe_filename(
            upload.filename or "attachment"
        )

        directory = (
            self._root
            / conversation_id
            / attachment_id
        )

        directory.mkdir(
            parents=True,
            exist_ok=True,
        )

        target = directory / filename

        sha = hashlib.sha256()

        total = 0

        try:
            with target.open("wb") as output:
                while True:
                    chunk = await upload.read(
                        1024 * 1024
                    )

                    if not chunk:
                        break

                    total += len(chunk)

                    if total > self._max_bytes:
                        raise AttachmentError(
                            "Attachment exceeds configured "
                            "upload size limit"
                        )

                    sha.update(chunk)
                    output.write(chunk)

        except Exception:
            target.unlink(
                missing_ok=True
            )

            raise

        finally:
            await upload.close()

        kind = _kind_for(filename)

        mime_type = (
            upload.content_type
            or mimetypes.guess_type(filename)[0]
        )

        relative = target.relative_to(
            self._root
        ).as_posix()

        virtual_path = f"/uploads/{relative}"

        extracted_virtual_path: str | None = None

        metadata: dict[str, object] = {
            "extension": target.suffix.lower(),
        }

        try:
            extracted = await self._extract(
                target,
                kind,
            )

            if extracted is not None:
                extracted_path = (
                    directory
                    / f"{filename}.extracted.txt"
                )

                extracted_path.write_text(
                    extracted,
                    encoding="utf-8",
                )

                extracted_relative = (
                    extracted_path.relative_to(
                        self._root
                    ).as_posix()
                )

                extracted_virtual_path = (
                    f"/uploads/"
                    f"{extracted_relative}"
                )

                metadata["extracted_chars"] = (
                    len(extracted)
                )

        except Exception as exc:
            # Original file is still useful to the agent.
            metadata["extraction_error"] = str(exc)

        return await self._repository.add_attachment(
            attachment_id=attachment_id,
            conversation_id=conversation_id,
            filename=filename,
            mime_type=mime_type,
            kind=kind,
            virtual_path=virtual_path,
            extracted_virtual_path=extracted_virtual_path,
            size_bytes=total,
            sha256=sha.hexdigest(),
            status=AttachmentStatus.READY,
            metadata=metadata,
        )

    async def _extract(
        self,
        path: Path,
        kind: AttachmentKind,
    ) -> str | None:
        if kind == AttachmentKind.TEXT:
            return await asyncio.to_thread(
                self._read_text,
                path,
            )

        if kind == AttachmentKind.PDF:
            return await asyncio.to_thread(
                self._read_pdf,
                path,
            )

        if kind == AttachmentKind.DOCX:
            return await asyncio.to_thread(
                self._read_docx,
                path,
            )

        return None

    def _clip(
        self,
        value: str,
    ) -> str:
        if len(value) <= self._max_extracted_chars:
            return value

        return value[
            : self._max_extracted_chars
        ]

    def _read_text(
        self,
        path: Path,
    ) -> str:
        raw = path.read_bytes()

        try:
            text = raw.decode("utf-8")

        except UnicodeDecodeError:
            text = raw.decode(
                "utf-8",
                errors="replace",
            )

        return self._clip(text)

    def _read_pdf(
        self,
        path: Path,
    ) -> str:
        reader = PdfReader(
            str(path)
        )

        sections: list[str] = []

        for index, page in enumerate(
            reader.pages,
            start=1,
        ):
            text = page.extract_text() or ""

            sections.append(
                f"# Page {index}\n\n{text}"
            )

        return self._clip(
            "\n\n".join(sections)
        )

    def _read_docx(
        self,
        path: Path,
    ) -> str:
        document = Document(
            str(path)
        )

        chunks: list[str] = []

        for paragraph in document.paragraphs:
            value = paragraph.text.strip()

            if value:
                chunks.append(value)

        for table in document.tables:
            for row in table.rows:
                values = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                chunks.append(
                    " | ".join(values)
                )

        return self._clip(
            "\n\n".join(chunks)
        )

"""Deterministic semantic verification for structured artifact formats.

Compares structured artifacts (PDF, DOCX, XLSX, CSV, ZIP, images) by their
meaningful content rather than by raw container bytes, so formatting noise
does not create false failures.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import zipfile

from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from server.src.skills.evaluation.fixtures import ReplayFixtureStore
from server.src.skills.representation.skill import (
    OutcomeAssertion,
    OutcomeAssertionType,
)


@dataclass(slots=True)
class FormatVerification:
    passed: bool
    score: float
    reason: str
    actual: Any = None
    expected: Any = None


class FormatVerifier:
    """Compare structured artifacts semantically rather than by raw bytes."""

    def __init__(self, fixtures: ReplayFixtureStore) -> None:
        self._fixtures = fixtures

    async def verify(
        self,
        assertion: OutcomeAssertion,
        actual_path: Path,
    ) -> FormatVerification:
        digest = assertion.expected_object_sha256
        if not digest:
            raise ValueError(
                f"{assertion.type.value} requires expected_object_sha256"
            )

        expected_bytes = await self._fixtures.read_object(digest)

        kind = assertion.type

        if kind == OutcomeAssertionType.PDF_SEMANTIC_EQUALS:
            return self._pdf(expected_bytes, actual_path, assertion.similarity_threshold)

        if kind == OutcomeAssertionType.DOCX_SEMANTIC_EQUALS:
            return self._docx(expected_bytes, actual_path, assertion.similarity_threshold)

        if kind == OutcomeAssertionType.XLSX_SEMANTIC_EQUALS:
            return self._xlsx(expected_bytes, actual_path, assertion.similarity_threshold)

        if kind == OutcomeAssertionType.CSV_SEMANTIC_EQUALS:
            return self._csv(expected_bytes, actual_path, assertion.similarity_threshold)

        if kind == OutcomeAssertionType.ZIP_SEMANTIC_EQUALS:
            return self._zip(expected_bytes, actual_path)

        if kind == OutcomeAssertionType.IMAGE_SEMANTIC_EQUALS:
            return self._image(expected_bytes, actual_path, assertion.similarity_threshold)

        raise ValueError(f"unsupported format assertion: {kind.value}")

    # ------------------------------------------------------------------
    # Shared
    # ------------------------------------------------------------------

    @staticmethod
    def _normalise_text(value: str) -> str:
        return "\n".join(
            line.rstrip()
            for line in value.replace("\r\n", "\n").replace("\r", "\n").strip().splitlines()
        )

    @classmethod
    def _text_similarity(cls, expected: str, actual: str) -> float:
        return SequenceMatcher(
            None,
            cls._normalise_text(expected),
            cls._normalise_text(actual),
            autojunk=False,
        ).ratio()

    @staticmethod
    def _json_safe(value: Any) -> Any:
        if value is None or isinstance(value, (str, int, float, bool)):
            return value
        if hasattr(value, "isoformat"):
            try:
                return value.isoformat()
            except Exception:
                pass
        return str(value)

    # ------------------------------------------------------------------
    # PDF
    # ------------------------------------------------------------------

    @classmethod
    def _pdf_signature(cls, source: bytes | Path) -> dict[str, Any]:
        from pypdf import PdfReader

        reader = PdfReader(
            io.BytesIO(source) if isinstance(source, bytes) else str(source)
        )

        pages = [
            cls._normalise_text(page.extract_text() or "")
            for page in reader.pages
        ]

        return {
            "page_count": len(pages),
            "pages": pages,
            "text": "\n\n".join(pages),
        }

    @classmethod
    def _pdf(
        cls,
        expected_bytes: bytes,
        actual_path: Path,
        threshold: float,
    ) -> FormatVerification:
        if not actual_path.is_file():
            return FormatVerification(False, 0.0, "PDF is missing")

        try:
            expected = cls._pdf_signature(expected_bytes)
            actual = cls._pdf_signature(actual_path)
        except Exception as exc:
            return FormatVerification(
                False,
                0.0,
                f"PDF parse failed: {type(exc).__name__}: {exc}",
            )

        page_score = 1.0 if actual["page_count"] == expected["page_count"] else 0.0
        text_score = cls._text_similarity(expected["text"], actual["text"])
        score = 0.25 * page_score + 0.75 * text_score
        passed = page_score == 1.0 and text_score >= threshold

        return FormatVerification(
            passed=passed,
            score=score,
            reason=f"PDF pages={actual['page_count']}, text_similarity={text_score:.3f}",
            actual={
                "page_count": actual["page_count"],
                "text_similarity": text_score,
            },
            expected={
                "page_count": expected["page_count"],
                "minimum_text_similarity": threshold,
            },
        )

    # ------------------------------------------------------------------
    # DOCX
    # ------------------------------------------------------------------

    @classmethod
    def _docx_signature(cls, source: bytes | Path) -> dict[str, Any]:
        from docx import Document

        document = Document(
            io.BytesIO(source) if isinstance(source, bytes) else str(source)
        )

        paragraphs = [
            cls._normalise_text(paragraph.text)
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        tables: list[list[list[str]]] = []
        for table in document.tables:
            tables.append(
                [
                    [cls._normalise_text(cell.text) for cell in row.cells]
                    for row in table.rows
                ]
            )

        signature_text = (
            "\n".join(paragraphs)
            + "\n"
            + json.dumps(tables, ensure_ascii=False, sort_keys=True)
        )

        return {
            "paragraphs": paragraphs,
            "tables": tables,
            "text": signature_text,
        }

    @classmethod
    def _docx(
        cls,
        expected_bytes: bytes,
        actual_path: Path,
        threshold: float,
    ) -> FormatVerification:
        if not actual_path.is_file():
            return FormatVerification(False, 0.0, "DOCX is missing")

        try:
            expected = cls._docx_signature(expected_bytes)
            actual = cls._docx_signature(actual_path)
        except Exception as exc:
            return FormatVerification(
                False,
                0.0,
                f"DOCX parse failed: {type(exc).__name__}: {exc}",
            )

        score = cls._text_similarity(expected["text"], actual["text"])

        return FormatVerification(
            passed=score >= threshold,
            score=score,
            reason=f"DOCX semantic similarity={score:.3f}",
            actual={
                "paragraphs": len(actual["paragraphs"]),
                "tables": len(actual["tables"]),
            },
            expected={
                "paragraphs": len(expected["paragraphs"]),
                "tables": len(expected["tables"]),
                "minimum_similarity": threshold,
            },
        )

    # ------------------------------------------------------------------
    # XLSX
    # ------------------------------------------------------------------

    @classmethod
    def _xlsx_signature(cls, source: bytes | Path) -> dict[str, Any]:
        from openpyxl import load_workbook

        workbook = load_workbook(
            io.BytesIO(source) if isinstance(source, bytes) else source,
            read_only=True,
            # Preserve formulas instead of depending on cached Excel
            # recalculation state.
            data_only=False,
        )

        sheets: dict[str, dict[str, Any]] = {}

        try:
            for sheet in workbook.worksheets:
                cells: dict[str, Any] = {}
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is None:
                            continue
                        cells[cell.coordinate] = cls._json_safe(cell.value)
                sheets[sheet.title] = cells
        finally:
            workbook.close()

        return {
            "sheet_names": list(sheets),
            "sheets": sheets,
        }

    @classmethod
    def _xlsx(
        cls,
        expected_bytes: bytes,
        actual_path: Path,
        threshold: float,
    ) -> FormatVerification:
        if not actual_path.is_file():
            return FormatVerification(False, 0.0, "XLSX is missing")

        try:
            expected = cls._xlsx_signature(expected_bytes)
            actual = cls._xlsx_signature(actual_path)
        except Exception as exc:
            return FormatVerification(
                False,
                0.0,
                f"XLSX parse failed: {type(exc).__name__}: {exc}",
            )

        expected_cells = {
            (sheet, coordinate): value
            for sheet, cells in expected["sheets"].items()
            for coordinate, value in cells.items()
        }

        actual_cells = {
            (sheet, coordinate): value
            for sheet, cells in actual["sheets"].items()
            for coordinate, value in cells.items()
        }

        keys = set(expected_cells) | set(actual_cells)

        if not keys:
            cell_score = 1.0
        else:
            matches = sum(
                1
                for key in keys
                if expected_cells.get(key) == actual_cells.get(key)
            )
            cell_score = matches / len(keys)

        sheet_score = 1.0 if expected["sheet_names"] == actual["sheet_names"] else 0.0
        score = 0.20 * sheet_score + 0.80 * cell_score
        passed = sheet_score == 1.0 and cell_score >= threshold

        return FormatVerification(
            passed=passed,
            score=score,
            reason=f"XLSX sheet_match={bool(sheet_score)}, cell_similarity={cell_score:.3f}",
            actual={
                "sheets": actual["sheet_names"],
                "nonempty_cells": len(actual_cells),
            },
            expected={
                "sheets": expected["sheet_names"],
                "nonempty_cells": len(expected_cells),
                "minimum_cell_similarity": threshold,
            },
        )

    # ------------------------------------------------------------------
    # CSV
    # ------------------------------------------------------------------

    @staticmethod
    def _csv_rows(source: bytes | Path) -> list[list[str]]:
        if isinstance(source, bytes):
            text = source.decode("utf-8-sig", errors="replace")
        else:
            text = source.read_text(encoding="utf-8-sig", errors="replace")

        return [
            [cell.strip() for cell in row]
            for row in csv.reader(io.StringIO(text))
        ]

    @classmethod
    def _csv(
        cls,
        expected_bytes: bytes,
        actual_path: Path,
        threshold: float,
    ) -> FormatVerification:
        if not actual_path.is_file():
            return FormatVerification(False, 0.0, "CSV is missing")

        expected = cls._csv_rows(expected_bytes)
        actual = cls._csv_rows(actual_path)

        expected_lines = [json.dumps(row, ensure_ascii=False) for row in expected]
        actual_lines = [json.dumps(row, ensure_ascii=False) for row in actual]

        score = SequenceMatcher(
            None,
            expected_lines,
            actual_lines,
            autojunk=False,
        ).ratio()

        return FormatVerification(
            passed=score >= threshold,
            score=score,
            reason=f"CSV row similarity={score:.3f}",
            actual={"rows": len(actual)},
            expected={"rows": len(expected), "minimum_similarity": threshold},
        )

    # ------------------------------------------------------------------
    # ZIP
    # ------------------------------------------------------------------

    @staticmethod
    def _zip_signature(source: bytes | Path) -> dict[str, str]:
        handle = io.BytesIO(source) if isinstance(source, bytes) else source

        with zipfile.ZipFile(handle, "r") as archive:
            result: dict[str, str] = {}
            for info in archive.infolist():
                if info.is_dir():
                    continue

                name = info.filename.replace("\\", "/")
                content = archive.read(info)
                result[name] = hashlib.sha256(content).hexdigest()

            return result

    @classmethod
    def _zip(
        cls,
        expected_bytes: bytes,
        actual_path: Path,
    ) -> FormatVerification:
        if not actual_path.is_file():
            return FormatVerification(False, 0.0, "ZIP is missing")

        try:
            expected = cls._zip_signature(expected_bytes)
            actual = cls._zip_signature(actual_path)
        except Exception as exc:
            return FormatVerification(
                False,
                0.0,
                f"ZIP parse failed: {type(exc).__name__}: {exc}",
            )

        expected_names = set(expected)
        actual_names = set(actual)
        union = expected_names | actual_names

        name_score = 1.0 if not union else len(expected_names & actual_names) / len(union)

        content_score = (
            1.0
            if not expected
            else sum(
                1
                for name, digest in expected.items()
                if actual.get(name) == digest
            ) / len(expected)
        )

        score = 0.40 * name_score + 0.60 * content_score
        passed = expected_names == actual_names and content_score == 1.0

        return FormatVerification(
            passed=passed,
            score=score,
            reason=f"ZIP member_match={name_score:.3f}, content_match={content_score:.3f}",
            actual={"members": sorted(actual)},
            expected={"members": sorted(expected)},
        )

    # ------------------------------------------------------------------
    # Images
    # ------------------------------------------------------------------

    @staticmethod
    def _average_hash(image: Any, hash_size: int = 16) -> int:
        from PIL import Image

        resampling = getattr(Image, "Resampling", Image).LANCZOS

        gray = (
            image.convert("L").resize((hash_size, hash_size), resampling)
        )

        pixels = list(gray.getdata())
        average = sum(pixels) / len(pixels)

        value = 0
        for pixel in pixels:
            value = (value << 1) | int(pixel >= average)

        return value

    @staticmethod
    def _mean_rgb(image: Any) -> list[int]:
        """Mean channel values (0-255), so flat-colour comparisons are colour-aware.

        A pure average-hash from a solid fill is all-ones regardless of the
        actual colour (every resampled pixel equals the mean), so two entirely
        different flat colours would otherwise compare as identical.
        """
        from PIL import Image

        resampling = getattr(Image, "Resampling", Image).LANCZOS

        small = image.convert("RGB").resize((8, 8), resampling)
        channels = small.split()

        return [
            round(sum(channel.getdata()) / len(list(channel.getdata())))
            for channel in channels
        ]

    @classmethod
    def _image_signature(cls, source: bytes | Path) -> dict[str, Any]:
        from PIL import Image

        handle = io.BytesIO(source) if isinstance(source, bytes) else source

        with Image.open(handle) as image:
            return {
                "width": image.width,
                "height": image.height,
                "format": image.format,
                "mode": image.mode,
                "ahash": cls._average_hash(image),
                "mean_rgb": cls._mean_rgb(image),
            }

    @classmethod
    def _color_similarity(
        cls,
        expected: list[int],
        actual: list[int],
    ) -> float:
        """Euclidean distance in RGB space normalised into a 0..1 similarity."""
        distance = sum(
            (expected_channel - actual_channel) ** 2
            for expected_channel, actual_channel in zip(expected, actual)
        ) ** 0.5
        max_distance = (3 * 255**2) ** 0.5
        return max(0.0, min(1.0, 1.0 - distance / max_distance))

    @classmethod
    def _image(
        cls,
        expected_bytes: bytes,
        actual_path: Path,
        threshold: float,
    ) -> FormatVerification:
        if not actual_path.is_file():
            return FormatVerification(False, 0.0, "image is missing")

        try:
            expected = cls._image_signature(expected_bytes)
            actual = cls._image_signature(actual_path)
        except Exception as exc:
            return FormatVerification(
                False,
                0.0,
                f"image parse failed: {type(exc).__name__}: {exc}",
            )

        bits = 16 * 16
        distance = (expected["ahash"] ^ actual["ahash"]).bit_count()
        visual_similarity = 1.0 - distance / bits

        color_similarity = cls._color_similarity(
            expected["mean_rgb"],
            actual["mean_rgb"],
        )

        dimension_match = (
            actual["width"] == expected["width"]
            and actual["height"] == expected["height"]
        )

        format_match = actual["format"] == expected["format"]

        score = (
            0.20 * float(dimension_match)
            + 0.10 * float(format_match)
            + 0.55 * visual_similarity
            + 0.15 * color_similarity
        )

        passed = (
            dimension_match
            and visual_similarity >= threshold
            and color_similarity >= threshold
        )

        return FormatVerification(
            passed=passed,
            score=score,
            reason=(
                f"image {actual['width']}x{actual['height']}, "
                f"perceptual_similarity={visual_similarity:.3f}, "
                f"color_similarity={color_similarity:.3f}"
            ),
            actual={
                "width": actual["width"],
                "height": actual["height"],
                "format": actual["format"],
                "mode": actual["mode"],
                "perceptual_similarity": visual_similarity,
                "color_similarity": color_similarity,
            },
            expected={
                "width": expected["width"],
                "height": expected["height"],
                "format": expected["format"],
                "mode": expected["mode"],
                "minimum_perceptual_similarity": threshold,
                "minimum_color_similarity": threshold,
            },
        )